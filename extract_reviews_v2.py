"""
Extract reviews from the 'Service & Artist Reviews' sheet bundled in the daily
bookings xlsx. Uses an EXACT join via Unique Order ID against the master
bookings data.json.
Rebuilt after an environment reset.
"""
import json, re, sys, os, csv
from datetime import datetime
import openpyxl

BOOKINGS_XLSX = sys.argv[1]
DATA_JSON = sys.argv[2] if len(sys.argv) > 2 else 'data.json'
REVIEWS_JSON = sys.argv[3] if len(sys.argv) > 3 else 'reviews.json'

def norm_provider(s):
    return re.sub(r'\s+', ' ', (s or '').strip()).lower()

SP_CATEGORY_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'sp_category.csv')
sp_model_lookup = {}
if os.path.exists(SP_CATEGORY_PATH):
    with open(SP_CATEGORY_PATH, newline='', encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            name = r.get('Smile Providers:') or ''
            cat = r.get('Category') or ''
            if name.strip():
                sp_model_lookup[norm_provider(name)] = cat.strip()

bookings_data = json.load(open(DATA_JSON))
orders_by_id = {o['id']: o for o in bookings_data['orders']}

# Build category lookup from existing lineItems (prefix-based, handles compound service names)
CATEGORY_LOOKUP = {}
for li in bookings_data['lineItems']:
    svc = li['service']
    sku = li['sku']
    prefix = svc + ' – '
    variant = sku[len(prefix):] if sku.startswith(prefix) else ''
    CATEGORY_LOOKUP[(svc, variant)] = (li['category'], li['group'])

SVC_ONLY_FALLBACK = {}
for (svc, variant), (cat, grp) in CATEGORY_LOOKUP.items():
    SVC_ONLY_FALLBACK.setdefault(svc, (cat, grp))

def categorize_service(svc_raw, variant_raw):
    if variant_raw == svc_raw:
        variant_raw = ''
    key = (svc_raw, variant_raw)
    if key in CATEGORY_LOOKUP:
        return CATEGORY_LOOKUP[key]
    if svc_raw in SVC_ONLY_FALLBACK:
        return SVC_ONLY_FALLBACK[svc_raw]
    return (svc_raw or 'Unknown', svc_raw or 'Unknown')

wb = openpyxl.load_workbook(BOOKINGS_XLSX, data_only=True)
if 'Service & Artist Reviews' not in wb.sheetnames:
    print("No 'Service & Artist Reviews' sheet in this file — nothing to do.")
    sys.exit(0)

ws = wb['Service & Artist Reviews']
idx = {c.value: i for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}

def to_iso_date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, str) and v:
        return v[:10]
    return None

new_reviews = []
new_service_ratings = []
unmatched_orders = set()
sheet_dates = []

for row in ws.iter_rows(min_row=2, values_only=True):
    order_id = row[idx['Unique Order ID']]
    if not order_id:
        continue
    review_type = row[idx['Review Type']]
    rating = row[idx['Rating']]
    if rating is None:
        continue
    order = orders_by_id.get(order_id)
    if not order:
        unmatched_orders.add(order_id)
        continue
    artist = order['provider']
    sched_date = order['scheduled']
    if sched_date:
        sheet_dates.append(sched_date)
    model = sp_model_lookup.get(norm_provider(artist), 'Churned / Unmapped') if artist else 'Churned / Unmapped'
    comment = row[idx['Comment']] or ''
    tags = row[idx['Tags']] or ''
    customer = row[idx['Customer Name']] or ''

    if review_type == 'EMPLOYEE':
        new_reviews.append({
            'orderId': order_id, 'source': 'exact', 'customer': customer,
            'artist': artist, 'rating': int(rating), 'date': sched_date,
            'spModel': model, 'comment': comment, 'tags': tags,
        })
    elif review_type == 'SERVICE':
        svc_raw = (row[idx['Reviewed Service']] or '').strip()
        if not svc_raw:
            continue
        # Reviewed Service format: "Service Name - Variant" or just "Service Name"
        if ' - ' in svc_raw:
            idx_dash = svc_raw.rfind(' - ')
            svc_name, svc_variant = svc_raw[:idx_dash].strip(), svc_raw[idx_dash+3:].strip()
        else:
            svc_name, svc_variant = svc_raw, ''
        category, service_label = categorize_service(svc_name, svc_variant)
        new_service_ratings.append({
            'orderId': order_id, 'source': 'exact', 'artist': artist, 'spModel': model,
            'customer': customer, 'date': sched_date, 'service': svc_name,
            'group': service_label, 'category': category, 'rating': int(rating),
            'comment': comment, 'tags': tags,
        })

sheet_min_date = min(sheet_dates) if sheet_dates else None
sheet_max_date = max(sheet_dates) if sheet_dates else None

existing = json.load(open(REVIEWS_JSON)) if os.path.exists(REVIEWS_JSON) else {'reviews': [], 'serviceRatings': []}

def keep_old(entry):
    if entry.get('orderId'):
        return False
    d = entry.get('date')
    return not (d and sheet_min_date and d >= sheet_min_date)

kept_reviews = [r for r in existing.get('reviews', []) if keep_old(r) and not r.get('orderId')]
kept_service_ratings = [r for r in existing.get('serviceRatings', []) if keep_old(r) and not r.get('orderId')]

prev_exact_reviews = {r['orderId']: r for r in existing.get('reviews', []) if r.get('orderId')}
prev_exact_service = [r for r in existing.get('serviceRatings', []) if r.get('orderId')]

touched_order_ids = set(r['orderId'] for r in new_reviews) | set(r['orderId'] for r in new_service_ratings)
carried_reviews = [r for oid, r in prev_exact_reviews.items() if oid not in touched_order_ids]
carried_service = [r for r in prev_exact_service if r['orderId'] not in touched_order_ids]

final_reviews = kept_reviews + carried_reviews + new_reviews
final_service_ratings = kept_service_ratings + carried_service + new_service_ratings

json.dump({'reviews': final_reviews, 'serviceRatings': final_service_ratings}, open(REVIEWS_JSON, 'w'))

print(f"New sheet covers: {sheet_min_date} -> {sheet_max_date}")
print(f"New EMPLOYEE (artist) reviews: {len(new_reviews)}")
print(f"New SERVICE ratings: {len(new_service_ratings)}")
print(f"Unmatched order IDs (not found in bookings data): {len(unmatched_orders)}")
if unmatched_orders:
    print("  ", sorted(unmatched_orders)[:10])
print(f"Legacy (pre-{sheet_min_date}) reviews kept: {len(kept_reviews)}")
print(f"Legacy service ratings kept: {len(kept_service_ratings)}")
print(f"Total reviews.json: {len(final_reviews)} reviews, {len(final_service_ratings)} service ratings")
